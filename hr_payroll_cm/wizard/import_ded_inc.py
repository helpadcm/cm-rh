from odoo import models, fields, _, api
from babel.dates import format_date
import calendar
from datetime import datetime, timedelta
from odoo.exceptions import ValidationError
from openpyxl import load_workbook
import base64
from io import BytesIO

class importIncomesDeductions(models.TransientModel):
    _name = "import.dedutions.incomes"
    _description = "Importar ingresos y deduducciones"

    options = fields.Selection([('incomes','Ingresos'),('deductions','Deducciones')],string="Opcion",default="deductions")
    format_doc = fields.Binary(string="Formato de importacion")
    doc_name = fields.Char(string="Nombre de Documento")

    def import_document(self):
        # print ("///////////////////////////")
        file_content = base64.b64decode(self.format_doc)

        wb = load_workbook(filename=BytesIO(file_content))
        ws = wb.active  # toma la primera hoja

        # Iterar por filas
        for row in ws.iter_rows(min_row=2, values_only=True):  # min_row=2 para saltar encabezados
            employee_name, description, ded_type, initial_date, final_date, monthly_amount, amount_total, fees = row

            employee_id = False
            type_id = False
            fees_values = False
            fees_qty = 0

            if not employee_name:
                raise ValidationError("Debe agregar nombre del empleado")
            else:
                employee_id = self.env['hr.employee'].search([('name','=',employee_name)])
                if not employee_id:
                    raise ValidationError(f"""No existe empleado con el nombre {employee_name}""")

            if not ded_type:
                raise ValidationError("Debe agregar el tipo de deduccion o ingreso")
            else:
                type_id = self.env['hr.salary.attachment.type'].search([('name','=',ded_type)])
                if not type_id:
                    raise ValidationError(f"""No existe tipo de deduccion o ingreso con el nombre {ded_type}""")

            if not initial_date:
                raise ValidationError("Debe agregar la fecha de inicio")

            if not description:
                description = type_id.name

            if not monthly_amount or monthly_amount <= 0:
                raise ValidationError("Debe ingresar el monto mensual o debe ser mayor que 0")

            if not amount_total or amount_total <= 0:
                raise ValidationError("Debe ingresar el monto total o debe ser mayor que 0")

            if fees:
                fees_values = True
                fees_qty = fees

            deduction_id = self.env['hr.salary.attachment'].create({
                'employee_ids': [(4, employee_id.id)],
                'deduction_type_id': type_id.id,
                'description': description,
                'date_start': initial_date,
                'date_end': final_date,
                'monthly_amount': monthly_amount,
                'total_amount': amount_total,
                'by_quotes': fees_values,
                'quotes_number': fees_qty
            })

            if fees_values:
                deduction_id.create_plan()

            # print (deduction_id)
        # print (a)
        # return True